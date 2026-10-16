"""通用工具：脱敏、导入导出、备份"""
import json
import shutil
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.core.config import settings


def mask_id_card(id_card: str | None) -> str | None:
    if not id_card:
        return id_card
    if len(id_card) <= 8:
        return id_card[:2] + "*" * (len(id_card) - 2)
    return id_card[:4] + "*" * (len(id_card) - 8) + id_card[-4:]


def mask_phone(phone: str | None) -> str | None:
    if not phone:
        return phone
    if len(phone) >= 7:
        return phone[:3] + "*" * (len(phone) - 7) + phone[-4:]
    return phone[:1] + "*" * (len(phone) - 1)


def serialize_student(s, mask: bool = False) -> dict:
    """将学生 ORM 对象序列化为字典（可选脱敏）"""
    return {
        "id": s.id,
        "student_no": s.student_no,
        "name": s.name,
        "gender": s.gender,
        "birth_date": s.birth_date.isoformat() if s.birth_date else None,
        "id_card": mask_id_card(s.id_card) if mask else s.id_card,
        "nation": s.nation,
        "political": s.political,
        "native_place": s.native_place,
        "origin": s.origin,
        "phone": mask_phone(s.phone) if mask else s.phone,
        "address": s.address,
        "photo_url": s.photo_url,
        "emergency_name": s.emergency_name,
        "emergency_phone": mask_phone(s.emergency_phone) if mask else s.emergency_phone,
        "family_note": s.family_note,
        "class_id": s.class_id,
        "class_name": s.class_info.class_name if s.class_info else None,
        "family_contacts": [
            {
                "id": f.id,
                "parent_name": f.parent_name,
                "relation": f.relation,
                "phone": mask_phone(f.phone) if mask else f.phone,
                "live_together": f.live_together,
                "economy_level": f.economy_level,
            }
            for f in (s.family_contacts or [])
        ],
    }


def export_rows_to_xlsx(headers: list[str], rows: list[list], sheet_name: str = "Sheet1") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_xlsx(content: bytes) -> list[dict]:
    """解析 xlsx 首行为表头，返回 [{列名: 值}]"""
    wb = load_workbook(BytesIO(content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    result = []
    for r in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        result.append({headers[i]: r[i] if i < len(r) else None for i in range(len(headers))})
    return result


def do_backup() -> str:
    """复制数据库文件到备份目录，返回备份文件名；保留最近 N 份"""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = settings.backup_dir / f"class_manager_{ts}.db"
    if settings.db_path.exists():
        shutil.copy2(settings.db_path, dest)
    # 清理旧备份
    backups = sorted(settings.backup_dir.glob("class_manager_*.db"))
    if len(backups) > settings.backup_keep_count:
        for old in backups[: len(backups) - settings.backup_keep_count]:
            old.unlink(missing_ok=True)
    return dest.name


def list_backups() -> list[dict]:
    backups = sorted(
        settings.backup_dir.glob("class_manager_*.db"), reverse=True
    )
    return [
        {"name": b.name, "size": b.stat().st_size, "mtime": b.stat().st_mtime}
        for b in backups
    ]
